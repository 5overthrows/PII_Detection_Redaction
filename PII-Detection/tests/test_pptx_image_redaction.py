from __future__ import annotations

import io
import sys
import tempfile
from types import SimpleNamespace
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from backend import file_redaction


class PptxImageRedactionTests(unittest.TestCase):
    def test_redacts_presentation_text_and_embedded_image(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "input.pptx"
            output_path = Path(temp_dir) / "redacted.pptx"
            image_buffer = io.BytesIO()
            Image.new("RGB", (40, 30), "white").save(image_buffer, format="PNG")

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            text_box = slide.shapes.add_textbox(left=0, top=0, width=300, height=50)
            text_box.text_frame.text = "Contact: Alice Example"
            slide.shapes.add_picture(
                io.BytesIO(image_buffer.getvalue()), left=0, top=0, width=40, height=30
            )
            presentation.save(input_path)

            def analyze_slide_text(text: str) -> list[SimpleNamespace]:
                if "Alice Example" not in text:
                    return []
                return [SimpleNamespace(start=9, end=22, entity_type="PERSON")]

            def analyze_ocr_text(text: str) -> list[SimpleNamespace]:
                if "Jane Doe" not in text:
                    return []
                return [SimpleNamespace(start=0, end=8)]

            ocr_lines = [
                file_redaction.OCRLine(
                    text="Jane Doe",
                    polygon=((10, 10), (30, 10), (30, 20), (10, 20)),
                )
            ]
            with (
                patch("backend.pptx_redaction.analyze", side_effect=analyze_slide_text),
                patch.object(file_redaction, "_ocr_lines", return_value=ocr_lines),
                patch.object(file_redaction, "_analyze", side_effect=analyze_ocr_text),
            ):
                file_redaction.redact_file(input_path, output_path)

            redacted_presentation = Presentation(output_path)
            redacted_text = "\n".join(
                shape.text
                for shape in redacted_presentation.slides[0].shapes
                if shape.has_text_frame
            )
            self.assertIn("Contact: <PERSON>", redacted_text)
            self.assertNotIn("Alice Example", redacted_text)
            picture = next(
                shape
                for shape in redacted_presentation.slides[0].shapes
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            )
            with Image.open(io.BytesIO(picture.image.blob)) as redacted_image:
                self.assertEqual(redacted_image.convert("RGB").getpixel((20, 15)), (0, 0, 0))
                self.assertEqual(redacted_image.convert("RGB").getpixel((0, 0)), (255, 255, 255))


if __name__ == "__main__":
    unittest.main()
