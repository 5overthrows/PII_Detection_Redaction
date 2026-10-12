"""
PPTX extractor preserving presentation structure, slides, shapes, bullets, and tables.
"""
import os
from typing import List, Dict, Any
from pptx import Presentation
from extractors.models import DocumentContent, ContentBlock

class PptxExtractor:
    @staticmethod
    def extract(file_path: str) -> DocumentContent:
        prs = Presentation(file_path)
        file_name = os.path.basename(file_path)
        blocks: List[ContentBlock] = []
        raw_text_parts: List[str] = []
        block_counter = 0
        image_count = 0
        unsupported_count = 0

        for slide_idx, slide in enumerate(prs.slides, start=1):
            # Check for slide title
            title_text = ""
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title_text = slide.shapes.title.text_frame.text.strip()
                if title_text:
                    block = ContentBlock(
                        block_id=f"pptx_s{slide_idx}_title",
                        block_type="slide_title",
                        text=title_text,
                        location={"slide": slide_idx, "shape_type": "title"},
                        style_metadata={"font_bold": True}
                    )
                    blocks.append(block)
                    raw_text_parts.append(title_text)
                    block_counter += 1

            # Iterate shapes
            def walk_shapes(shapes, parent_path=""):
                nonlocal block_counter, image_count
                for shape_idx, shape in enumerate(shapes):
                    path = f"{parent_path}{shape_idx}"
                    if shape.shape_type == 13:  # PICTURE
                        image_count += 1
                    if getattr(shape, "has_chart", False):
                        unsupported_count += 1
                    if shape.shape_type == 6:  # GROUP
                        walk_shapes(shape.shapes, path + ".")
                        continue
                    # If shape is title and already handled, skip
                    if parent_path == "" and shape == slide.shapes.title:
                        continue

                    # Table handling
                    if shape.has_table:
                        table = shape.table
                        for r_idx, row in enumerate(table.rows):
                            row_cells_text = [cell.text.strip() for cell in row.cells]
                            row_str = " | ".join(row_cells_text)
                            blocks.append(ContentBlock(
                                block_id=f"pptx_s{slide_idx}_{path}_tbl_r{r_idx}", block_type="table_row",
                                text=row_str, location={"slide": slide_idx, "shape_path": path, "row": r_idx, "cols": len(row.cells)},
                                style_metadata={"is_header": (r_idx == 0)}
                            ))
                            raw_text_parts.append(row_str)
                            block_counter += 1
                    elif shape.has_text_frame:
                        for p_idx, p in enumerate(shape.text_frame.paragraphs):
                            p_text = p.text.strip()
                            if not p_text:
                                continue
                            is_bullet = p.level > 0 or p_text.startswith(("-", "•", "*"))
                            block = ContentBlock(
                                block_id=f"pptx_s{slide_idx}_sh{path}_p{p_idx}",
                                block_type="bullet" if is_bullet else "paragraph", text=p_text,
                                location={"slide": slide_idx, "shape": path, "paragraph": p_idx},
                                style_metadata={"level": p.level}
                            )
                            blocks.append(block)
                            raw_text_parts.append(p_text)
                            block_counter += 1

            walk_shapes(slide.shapes)

            # Speaker Notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                note_text = slide.notes_slide.notes_text_frame.text.strip()
                if note_text:
                    block = ContentBlock(
                        block_id=f"pptx_s{slide_idx}_notes",
                        block_type="metadata",
                        text=f"Speaker Notes: {note_text}",
                        location={"slide": slide_idx, "type": "notes"}
                    )
                    blocks.append(block)
                    raw_text_parts.append(note_text)
                    block_counter += 1

        full_raw_text = "\n".join(raw_text_parts)
        metadata = {
            "slide_count": len(prs.slides),
            "source_type": "pptx",
            "coverage": {
                "text_extracted": bool(full_raw_text.strip()),
                "ocr_performed": False,
                "image_count": image_count,
                "unsupported_content_count": unsupported_count,
                "warnings": ([f"{image_count} embedded image(s) were not OCR-scanned."] if image_count else []) +
                    ([f"{unsupported_count} chart(s) are not included in text extraction."] if unsupported_count else [])
            }
        }

        return DocumentContent(
            file_name=file_name,
            file_type="pptx",
            raw_text=full_raw_text,
            blocks=blocks,
            metadata=metadata
        )
